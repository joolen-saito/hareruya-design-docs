"""canonical model(JSON) から一時 .xlsx を書き戻す(仕様§4「再構成」)。

ここが本ツール内で唯一 openpyxl を書き込みに使う箇所である。読み取り側
(抽出・source検証)は一切openpyxlを使わない。再構成後は `ooxml`/`canonical`
の直接XML読取で再読込し、同じcanonical modelになるかを比較する
(比較は `verify.py`)。

styleId・relationship ID・XML属性順の一致は要求しない(仕様§4)。
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter


def _font_from_normalized(nf: dict) -> Font:
    color = nf.get("color")
    color_hex = None
    if color and color.get("hex"):
        color_hex = color["hex"].lstrip("#")
    return Font(
        bold=nf.get("bold", False),
        strike=nf.get("strike", False),
        italic=nf.get("italic", False),
        underline="single" if nf.get("underline") else None,
        size=nf.get("size"),
        name=nf.get("name"),
        color=color_hex,
    )


def _fill_from_normalized(nf: dict) -> Optional[PatternFill]:
    pattern = nf.get("pattern_type")
    if not pattern:
        return None
    fg = nf.get("fg_color") or {}
    bg = nf.get("bg_color") or {}
    fg_hex = fg.get("hex", "").lstrip("#") if fg.get("hex") else None
    bg_hex = bg.get("hex", "").lstrip("#") if bg.get("hex") else None
    kwargs: dict = {"patternType": pattern}
    if fg_hex:
        kwargs["fgColor"] = fg_hex
    if bg_hex:
        kwargs["bgColor"] = bg_hex
    return PatternFill(**kwargs)


def _side_from_normalized(side: dict) -> Side:
    style = side.get("style")
    if not style:
        return Side()
    color = side.get("color") or {}
    color_hex = color.get("hex", "").lstrip("#") if color.get("hex") else None
    return Side(style=style, color=color_hex)


def _border_from_normalized(nb: dict) -> Border:
    return Border(
        left=_side_from_normalized(nb.get("left", {})),
        right=_side_from_normalized(nb.get("right", {})),
        top=_side_from_normalized(nb.get("top", {})),
        bottom=_side_from_normalized(nb.get("bottom", {})),
        diagonal=_side_from_normalized(nb.get("diagonal", {})),
        diagonalUp=nb.get("diagonal_up", False),
        diagonalDown=nb.get("diagonal_down", False),
    )


def _alignment_from_normalized(na: dict) -> Alignment:
    return Alignment(
        horizontal=na.get("horizontal"),
        vertical=na.get("vertical"),
        wrapText=na.get("wrap_text", False),
        textRotation=na.get("text_rotation") or 0,
        indent=na.get("indent") or 0,
        shrinkToFit=na.get("shrink_to_fit", False),
    )


def _protection_from_normalized(np_: dict) -> Protection:
    return Protection(locked=np_.get("locked", True), hidden=np_.get("hidden", False))


def _inline_font(run: dict) -> InlineFont:
    return InlineFont(
        b=run.get("explicit_bold"),
        strike=run.get("explicit_strike"),
        i=run.get("explicit_italic"),
        u=("single" if run.get("explicit_underline") else None),
    )


def build_temp_workbook(model: dict, output_path: Path) -> Path:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    sheet_name = (model["source"]["sheet_name"] or "block")[:31]
    ws = wb.create_sheet(title=sheet_name)

    styles_table = model.get("styles", {})
    style_cache: dict[str, dict] = {}

    def styled_parts(fingerprint: str) -> dict:
        if fingerprint not in style_cache:
            nx = styles_table[fingerprint]["normalized_xf"]
            style_cache[fingerprint] = {
                "font": _font_from_normalized(nx["font"]),
                "fill": _fill_from_normalized(nx["fill"]),
                "border": _border_from_normalized(nx["border"]),
                "alignment": _alignment_from_normalized(nx["alignment"]),
                "protection": _protection_from_normalized(nx["protection"]),
                "number_format": nx.get("num_fmt_code") or "General",
                "quote_prefix": nx.get("quote_prefix", False),
            }
        return style_cache[fingerprint]

    block = model["block"]
    row_start, row_end = block["row_start"], block["row_end"]
    col_start, col_end = block["col_start"], block["col_end"]

    # Merge FIRST, then write. openpyxl.Worksheet.merge_cells() replaces every
    # non-anchor cell in the range with a fresh MergedCell and DISCARDS any
    # value/style previously set on it (observed empirically) - writing cell
    # content before merging silently loses the style of covered cells.
    # MergedCell also rejects `.value =` assignment (AttributeError), so
    # covered cells only ever receive style, never a value/formula.
    covered_addresses: set[str] = set()
    for ref in model.get("merged_ranges", []):
        min_row, min_col, max_row, max_col = _parse_ref(ref)
        ws.merge_cells(
            start_row=min_row - row_start + 1,
            start_column=min_col - col_start + 1,
            end_row=max_row - row_start + 1,
            end_column=max_col - col_start + 1,
        )
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if (r, c) == (min_row, min_col):
                    continue  # anchor keeps its regular Cell
                covered_addresses.add(f"{get_column_letter(c)}{r}")

    cells = model["cells"]
    for row_num in range(row_start, row_end + 1):
        for col_num in range(col_start, col_end + 1):
            addr = f"{get_column_letter(col_num)}{row_num}"
            cell_model = cells.get(addr)
            if cell_model is None or not cell_model.get("present_in_xml"):
                continue
            ws_cell = ws.cell(row=row_num - row_start + 1, column=col_num - col_start + 1)
            _apply_cell(ws_cell, cell_model, styled_parts, write_value=addr not in covered_addresses)

    rows_meta = model.get("rows", {})
    for row_num in range(row_start, row_end + 1):
        meta = rows_meta.get(str(row_num))
        if not meta or not meta.get("present_in_xml"):
            continue
        rd = ws.row_dimensions[row_num - row_start + 1]
        if meta.get("height") is not None:
            rd.height = meta["height"]
        rd.hidden = bool(meta.get("hidden"))
        rd.outlineLevel = meta.get("outline_level", 0) or 0
        # openpyxl.RowDimension.customHeight is a READ-ONLY computed property
        # (`return self.ht is not None`) - there is no public API to set it
        # independently of height. A row with an explicit calculated height
        # but custom_height=False in the original cannot be reconstructed
        # 1:1 through openpyxl; verify.py excludes custom_height from the
        # reconstruct-mode comparison specifically for this reason (documented
        # there, not silently dropped).

    cols_meta = model.get("columns", {})
    for col_num in range(col_start, col_end + 1):
        letters_abs = get_column_letter(col_num)
        meta = cols_meta.get(letters_abs)
        if not meta:
            continue
        letters_rel = get_column_letter(col_num - col_start + 1)
        cd = ws.column_dimensions[letters_rel]
        if meta.get("width") is not None:
            cd.width = meta["width"]
        cd.hidden = bool(meta.get("hidden"))
        cd.outlineLevel = meta.get("outline_level", 0) or 0
        # openpyxl.ColumnDimension.customWidth is likewise a read-only computed
        # property (`return bool(self.width)`) - see the customHeight note above.

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _parse_ref(ref: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return min_row, min_col, max_row, max_col


def _apply_cell(ws_cell, cell_model: dict, styled_parts, write_value: bool = True) -> None:
    style_info = cell_model.get("style")
    parts = styled_parts(style_info["fingerprint"]) if style_info else None

    cell_type = cell_model["type"]
    formula = cell_model.get("formula")
    value_lexical = cell_model.get("value_lexical")
    runs = cell_model.get("rich_text_runs") or []

    # `write_value=False` for merged-covered cells: openpyxl.MergedCell
    # rejects `.value =` assignment outright (AttributeError), and Excel
    # itself never stores real content in covered cells anyway - only the
    # anchor cell's value is meaningful. Style still applies to every cell.
    if write_value:
        if formula:
            ws_cell.value = f"={formula}"
        elif cell_type in ("s", "inlineStr", "str"):
            if len(runs) > 1:
                blocks = []
                for r in runs:
                    text = r["text"]
                    if not text:
                        continue
                    blocks.append(TextBlock(_inline_font(r), text))
                ws_cell.value = CellRichText(*blocks) if blocks else ""
            else:
                ws_cell.value = value_lexical or ""
        elif cell_type == "n":
            if value_lexical is not None:
                try:
                    num = float(value_lexical)
                    ws_cell.value = int(num) if num.is_integer() else num
                except ValueError:
                    ws_cell.value = value_lexical
        elif cell_type == "b":
            ws_cell.value = value_lexical == "1"
        elif cell_type == "e":
            # openpyxl has no first-class error literal writer; preserve lexically.
            ws_cell.value = value_lexical
        elif cell_type == "blank":
            pass
        else:
            ws_cell.value = value_lexical

    if parts:
        ws_cell.font = parts["font"]
        if parts["fill"] is not None:
            ws_cell.fill = parts["fill"]
        ws_cell.border = parts["border"]
        ws_cell.alignment = parts["alignment"]
        ws_cell.protection = parts["protection"]
        ws_cell.number_format = parts["number_format"]
        if parts["quote_prefix"]:
            ws_cell.quotePrefix = True
